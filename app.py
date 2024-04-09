import os 
from flask import *
from flask_bcrypt import Bcrypt
from flask_pymongo import PyMongo
from werkzeug.utils import secure_filename
import re
from docx import Document
from docxtpl import DocxTemplate
import tempfile
from pymongo import MongoClient

# root
# mongodb+srv://root:<password>@cluster0.0jmtyiz.mongodb.net/?retryWrites=true&w=majority&appName=Cluster0

# ^ Here we initialize the python flask app
app = Flask(__name__)
app.secret_key = 'jayganesh'
# app.config['MONGO_URI'] = 'mongodb://localhost:27017/pymongo'
MONGO_URI = "mongodb+srv://harsh:h1h2h3h4h5h6h7@cluster0.0jmtyiz.mongodb.net/pymongo?retryWrites=true&w=majority"
client = MongoClient(MONGO_URI)
db = client.get_database("pymongo")

# * Define the upload directory and allowed file extensions
UPLOAD_FOLDER = 'static'
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif'}
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
# mongo = PyMongo(app)
bcrypt = Bcrypt(app)

#~ Function 
@app.after_request
def add_no_cache_headers(response):
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
    response.headers['Pragma'] = 'no-cache'
    response.headers['Expires'] = '-1'
    return response

# ^ This is the route for `login`
@app.route('/login', methods=['GET', 'POST'])
def login():
    error = None
    if request.method == 'POST':
        users = db.users
        username = request.form['username']
        password = request.form['password']

        user = users.find_one({'username': username})
        if user and bcrypt.check_password_hash(user['password'], password):
            error = False
            session['username'] = username  # Storing username in session
            return redirect(url_for('index'))  # Redirect to dashboard upon successful login
        else:
            error = True
            return render_template('login.html', message='Invalid username or password', error=error)
        
    if 'username' in session:
        return redirect(url_for('index'))

    return render_template('login.html')

@app.route('/signup', methods=['GET', 'POST'])
def signup():
    error = None 
    if request.method == 'POST':
        users = db.users
        username = request.form['username']
        useremail= request.form['username']
        password = request.form['password']

        existing_user = users.find_one({'username': username})
        if existing_user:
            error = True
            return render_template('signup.html', message='Invalid username or password', error=error)

        hashed_password = bcrypt.generate_password_hash(password).decode('utf-8')
        users.insert_one({'username': username, 'useremail': useremail, 'password': hashed_password})
        error = False

        return render_template('login.html', message='Signup successful. You can now login.', error=error)

    return render_template('signup.html')
    
@app.route('/logout')
def logout():
    session.pop('username', None)
    return redirect(url_for('login'))

# ! Do not remove this 💀
@app.route("/")
def index():
    if 'username' in session:
        username = session['username']
        documents = db.template_cards.find()
        return render_template('index.html', username=username, documents=documents)
    else:
        return redirect(url_for('login'))
    

#~ Function to check if the file extension is allowed
def allowed_file(filename):
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


@app.route('/profile')
def profile():
    try:
        username = session.get('username')
        print("the value of the user id is", username)

        starred_documents = db.stars.find({'user_id': username}, {'document_id': 1})

        documents = []
        for starred_doc in starred_documents:
            document_id = starred_doc['document_id']
            document = db.template_cards.find_one({'_id': ObjectId(document_id)})
            if document:
                documents.append(document)

        starred_count = db.stars.count_documents({'user_id': username})

        used_documents = db.template_cards.find({'used_by': username})

        user_documents = db.template_cards.find({'creator': username})

        return render_template('profilepage.html', username=session.get("username"), starred_documents=documents, starred_count=starred_count, used_documents=used_documents, user_documents=user_documents)

    except Exception as e:
        print(f"Error: {str(e)}")
        return jsonify({'success': False, 'error': str(e)})


@app.route("/otherprofile/<username>")
def otherprofile(username):
    try:
        person = db.template_cards.find({'creator': username})

        if person:
            user_documents = db.template_cards.find({'creator': username})

            used_documents = db.template_cards.find({'used_by': username})

            return render_template('otherprofilepage.html', username=username, used_documents=used_documents, user_documents=user_documents)
        else:
            return "User not found"  # Handle the case where the user does not exist
    except Exception as e:
        return str(e)


@app.route("/upload")
def upload():
    if 'username' in session:
        username = session['username']
        return render_template("upload.html", username=username)
    else:
        return redirect(url_for('login'))
    

@app.route("/submit_upload", methods=["POST"])
def submit_upload():
    if request.method == "POST":
        title = request.form["document_title"]
        description = request.form["desc"]
        docType = request.form["dt"]
        amount = request.form["amount"]
        isPaid = request.form["money"]
        document_file = request.files['document']
        files = request.files.getlist('images[]')

        creator = session.get("username")

        if 'images[]' not in request.files:
            flash('No file part')
            return redirect(request.url)
        
        def allowed_document(filename):
            return '.' in filename and \
                filename.rsplit('.', 1)[1].lower() in {'docx', 'pdf'}

        
        if document_file and allowed_document(document_file.filename):
            document_filename = secure_filename(document_file.filename)
            document_file.save(os.path.join(app.config['UPLOAD_FOLDER'], document_filename))
            document_url = os.path.join(app.config['UPLOAD_FOLDER'], document_filename)
        else:
            flash('Invalid document file format. Only .docx and .pdf files are allowed.')
            return redirect(request.url)


        image_urls = []
        for image in files:
            if image and allowed_file(image.filename):
                filename = secure_filename(image.filename)
                image.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))
                image_url = os.path.join(app.config['UPLOAD_FOLDER'], filename)
                image_urls.append(image_url)
            else:
                flash('Invalid file format')
                return redirect(request.url)
        
        db.template_cards.insert_one({
            "title": title,
            "description": description,
            "image_url": image_urls,
            "document": document_url,
            "creator": creator,
            "views": 0,
            "stars": 0,
            "uses": 0,
            "badge": docType,
            "price": amount,
            "is_paid": isPaid,
            "comments": []
        })
        
        return redirect(url_for("index")) 

    return "Invalid request"

@app.route('/use_this/<document>')
def use_this(document):
    document_path = get_document_path_from_database(document)
    return redirect(url_for('enter_data', document=document, document_path=document_path))


def extract_placeholders(docx_file):
    placeholders = set()
    doc = Document(docx_file)
    pattern = r'\{\{([^}]+)\}\}'
    for paragraph in doc.paragraphs:
        matches = re.findall(pattern, paragraph.text)
        placeholders.update(matches)
    return placeholders

@app.route('/enter_data')
def enter_data():
    document = request.args.get('document')
    document_path = request.args.get('document_path')

    placeholders = extract_placeholders(document_path)

    return render_template('enter_data.html', document=document, placeholders=placeholders)

def get_document_path_from_database(document_name):
    print("teh name of the document is see here", document_name)
    document = db.template_cards.find_one({"title": document_name})

    if document:
        document_path = document.get("document")
        print(document_path)
        return document_path
    else:
        print("Document does not exist")
        return None  

@app.route('/submit_data', methods=['POST'])
def submit_data():
    if request.method == 'POST':
        document_name = request.form["documentName"]
        document_path = get_document_path_from_database(document_name)

        placeholders = extract_placeholders(document_path)
        print("Placeholders from document:", placeholders)

        form_data_dict = {key.strip(): value for key, value in request.form.items()}
        print("Form Data:", form_data_dict)

        filled_document_path = fill_document(document_path, form_data_dict)

        return send_file(filled_document_path, as_attachment=True)

    return 'Form submitted successfully'

def fill_document(document_path, placeholders_values):
    doc = DocxTemplate(document_path)
    doc.render(placeholders_values)
    
    filled_document_path = tempfile.NamedTemporaryFile(suffix='.docx', delete=False).name
    doc.save(filled_document_path)
    
    return filled_document_path

# ================== some essential actions ==================
from bson import ObjectId

# ^ route for increamenting the star count
@app.route('/star/<document_id>', methods=['POST'])
def star_document(document_id):
    try:
        user_id = session.get('username')  # If user ID is stored in the session
        print("user id is", user_id, session)

        document = db.template_cards.find_one({'_id': ObjectId(document_id)})

        if document:
            existing_star = db.stars.find_one({'document_id': document_id, 'user_id': user_id})

            if existing_star:
                db.stars.delete_one({'document_id': document_id, 'user_id': user_id})
                db.template_cards.update_one(
                    {'_id': document['_id']},
                    {'$inc': {'stars': -1}}
                )
                star_count = document['stars'] - 1
                return jsonify({'success': True, 'star_count': star_count, 'action': 'unstar'})
            else:
                db.stars.insert_one({'document_id': document_id, 'user_id': user_id})
                db.template_cards.update_one(
                    {'_id': document['_id']},
                    {'$inc': {'stars': 1}}
                )
                star_count = document['stars'] + 1
                return jsonify({'success': True, 'star_count': star_count, 'action': 'star'})
        else:
            return jsonify({'success': False, 'error': 'Document not found'})

    except Exception as e:
        print(f"Error: {str(e)}")
        return jsonify({'success': False, 'error': str(e)})
    
# ^ route for increamenting the used by count
from flask import jsonify, request

from flask import jsonify, request
from bson import ObjectId

@app.route('/use/<document_id>', methods=['POST'])
def use_document(document_id):
    try:
        username = session.get('username')  # If user ID is stored in the session
        print("user id is", session, username)

        db.template_cards.update_one(
            {'_id': ObjectId(document_id)},
            {'$inc': {'uses': 1}}
        )

        db.template_cards.update_one(
            {'_id': ObjectId(document_id)},
            {'$addToSet': {'used_by': username}}
        )

        db.document_usage.update_one(
            {'document_id': document_id},
            {'$inc': {'count': 1}},
            upsert=True  
        )

        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})

#^ To implement the search results this route is used...
from flask import request, render_template

from flask import request, render_template

@app.route('/search', methods=['GET'])
def search():
    search_query = request.args.get('q')

    if search_query:
        search_results = db.template_cards.find({'title': {'$regex': search_query, '$options': 'i'}})

        return render_template('index.html', search_results=search_results, search_query=search_query)
    else:
        return '<h1>Not Found!</h1>'
    
#^ to delete the document template
@app.route("/delete_document", methods=["POST"])
def delete_document():
    if request.method == 'POST':
        data = request.get_json()  # Assuming data contains the document ID to delete
        document_id = data.get("document_id")  # Assuming document_id is the key for the document ID
        print("data is", data, "document id is", document_id)
        if not document_id:
            return jsonify({"error": "Document ID not provided"}), 400

        result = db.template_cards.delete_one({"_id": ObjectId(document_id)})

        if result.deleted_count == 1:
            return jsonify({"success": True, "message": "Document template card deleted successfully"})
        else:
            return jsonify({"error": "Document template card not found"}), 404

#^ to update the document template
from flask import flash, redirect, render_template, request, url_for

@app.route("/update_document/<string:document_id>", methods=["GET", "POST"])
def update_document(document_id):
    if request.method == "GET":
        document = db.template_cards.find_one({"_id": ObjectId(document_id)})
        if not document:
            flash("Document not found")
            return redirect(url_for("index"))

        return render_template("update_document.html", document=document)
    
    elif request.method == "POST":
        title = request.form.get("document_title")
        document_type = request.form.get("dt")
        amount_price = request.form.get("amount")
        charge_status = request.form.get("money")
        description = request.form.get("desc")

        db.template_cards.update_one({"_id": ObjectId(document_id)}, {"$set": 
        {"title": title, "badge": document_type, "price": amount_price, "is_paid": charge_status, "description": description}})

        return redirect(url_for("index"))


if __name__ == '__main__':
    app.run(debug=True)
